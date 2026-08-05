import base64
import datetime
import hashlib
import json
import os
from functools import cached_property, partial
from io import BytesIO
from urllib.parse import quote

from ajax_helpers.mixins import AjaxHelpers, AjaxTaskMixin
from ajax_helpers.templatetags.ajax_helpers import button_javascript, post_json_js
from ajax_helpers.utils import ajax_command, is_ajax
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.http import Http404
from django.shortcuts import redirect
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils.html import escape
from django.utils.safestring import mark_safe
from django_datatables.columns import DateTimeColumn, DatatableColumn, ColumnLink, ColumnBase
from django_datatables.datatables import DatatableView
from django_datatables.helpers import row_button, DUMMY_ID
from django_menus.menu import AjaxMenuTemplateView, HtmlMenu, MenuItem, MenuItemDisplay, MenuMixin
from django_modals.decorators import ConfirmAjaxMethod
from django_modals.helper import reverse_modal
from openpyxl import Workbook

from cloud_backup.backup import Backup
from . import storage_setup
from .backup_local_files import BackupLocal, local_backup_path
from .config import config_index, config_names, safe_config, selected_config_name
from .sql_functions import get_schemas, get_schema_tables, get_table_column_names, get_table_data
from .tasks import ajax_backup
from .utils import allowed_to_restore, RESTORE_BLOCKED_MESSAGE


BADGE_COLOURS = {'Enabled': 'success', 'Disabled': 'danger', 'Suspended': 'warning'}


def status_table(rows):
    """protection_info/destination_status rows as one aligned table. The folder column is
    only rendered when something has a folder - versioning, object lock, soft delete and
    the Google Drive row are about the whole destination, not a key prefix."""
    folders = any(row.get('folder') for row in rows)
    folder_header = '<th>Folder</th>' if folders else ''
    html = ('<table class="table table-sm w-auto mb-0"><thead><tr><th>Check</th>'
            f'{folder_header}<th>Status</th><th>Detail</th></tr></thead><tbody>')
    for row in rows:
        colour = BADGE_COLOURS.get(row['status'], 'secondary')
        folder = (f'<td class="text-monospace text-nowrap">{escape(row.get("folder") or "")}</td>'
                  if folders else '')
        html += (f'<tr><td class="text-nowrap">{escape(row["label"])}</td>{folder}'
                 f'<td><span class="badge badge-{colour}">{escape(row["status"])}</span></td>'
                 f'<td class="text-muted">{escape(row.get("detail") or "")}</td></tr>')
    return html + '</tbody></table>'


def config_url(url_name, *url_args, config=None):
    """A backup page url that stays on the selected destination.

    The config rides in the query string rather than the path: the enhanced patterns end
    in a single-segment <str:schema> catch-all that a /<config>/ prefix would be
    ambiguous with, the url names branded projects reverse have to stay as they are, and
    ajax_helpers posts to window.location while datatables posts to
    window.location.search - so the tables, row buttons and storage info carry it for
    free. Modal urls cannot use it: show_modal appends ?modal_id=..., which would leave
    the url with two query strings - modals carry the config in their slug instead."""
    url = reverse(url_name, args=url_args)
    return f'{url}?config={quote(config)}' if config else url


class ConfigTab(MenuItem):
    """One destination in the tab strip. Which tab is current has to be set by the view:
    every database tab reverses the same url name and they differ only by ?config=,
    which neither of MenuItem.active's comparisons looks at."""

    def __init__(self, *args, active=False, **kwargs):
        self._active = active
        super().__init__(*args, link_type=MenuItem.HREF, **kwargs)

    @property
    def active(self):
        return self._active


class BackupConfigMixin:
    """Which BACKUP_CONFIGS destination a page is working on, and the tab strip that
    switches between them - one tab per destination plus Storage Setup, which is how that
    page is reached. Reads config_names()/safe_config() only - never Backup() - so it is
    safe on the storage setup page, whose job is describing configs that do not resolve."""

    # the setup page is one of the tabs rather than a destination, so it says so itself
    setup_page = False

    @cached_property
    def config_names(self):
        return config_names()

    @cached_property
    def config_name(self):
        return selected_config_name(self.request.GET.get('config'))

    @cached_property
    def config_index(self):
        return config_index(self.config_name)

    @property
    def multi_config(self):
        """A single-destination install keeps exactly the urls and markup it had before
        named configs existed: no tabs, and no ?config= on any link."""
        return len(self.config_names) > 1

    @property
    def url_config(self):
        return self.config_name if self.multi_config else None

    def page_item(self, url_name, text, *url_args, **kwargs):
        """Menu tuple for a link to another backup page, keeping the selected config."""
        return (config_url(url_name, *url_args, config=self.url_config), text,
                dict(kwargs, link_type=MenuItem.HREF))

    def config_home_url(self, name, config=None):
        """Where a config's tab goes: its database page, its file browser when the config
        has no database, or the setup page when the config does not resolve - that is the
        page that explains why."""
        config = config if config is not None else safe_config(name)
        # a single-destination install carries no ?config= at all
        name = name if self.multi_config else None
        if config is None:
            return config_url('cloud_backup:storage_setup', config=name)
        if config.include_db:
            return config_url('cloud_backup:backup_info', config=name)
        if len(config.dirs) == 1:
            return config_url('cloud_backup:backup_files', 0, config=name)
        return config_url('cloud_backup:backup_files_root', config=name)

    def unusable_config_redirect(self, request, require_db=False):
        """Send a page load for a destination this page cannot show to one that explains
        it: the setup page when the config's settings do not resolve, the file browser
        when the page needs a database and the config has it turned off. Page loads only -
        a stale table post must not be answered with a redirect."""
        if request.method != 'GET' or is_ajax(request):
            return None
        config = safe_config(self.config_name)
        if config is None or (require_db and not config.include_db):
            return redirect(self.config_home_url(self.config_name, config))
        return None

    def add_config_tabs(self):
        """One tab per destination, then Storage Setup. The strip is built even for a
        single destination, because the setup page has no other way in."""
        menu = self.add_menu('configs', menu_type='tabs')
        for name in self.config_names:
            config = safe_config(name)
            if config is None:
                icon = 'fas fa-exclamation-triangle text-danger'
            elif config.include_db:
                icon = 'fas fa-database'
            else:
                icon = 'fas fa-folder'
            # with one destination its name says nothing - the tab names the page instead
            menu.add_items(ConfigTab(self.config_home_url(name, config),
                                     MenuItemDisplay(name if self.multi_config else 'Backup',
                                                     font_awesome=icon),
                                     active=not self.setup_page and name == self.config_name))
        menu.add_items(ConfigTab(config_url('cloud_backup:storage_setup', config=self.url_config),
                                 MenuItemDisplay('Storage Setup', font_awesome='fas fa-cog'),
                                 active=self.setup_page))


def overwrite_visible_cell(table, row_no, column_name, html):
    """Like django_datatables' overwrite_cell, but hidden columns (e.g. '.id') render
    no <td> at all, so the td:nth-of-type position must be counted over visible
    columns only - overwrite_cell counts hidden ones and misses the cell."""
    visible = [c.column_name for c in table.columns if not c.options.get('hidden')]
    return ajax_command('html',
                        selector=f'#{table.table_id} #{row_no} td:nth-of-type({visible.index(column_name) + 1})',
                        html=html)


class TableBackup(BackupConfigMixin, AjaxTaskMixin, AjaxHelpers):

    tasks = {'backup': ajax_backup}

    @cached_property
    def backup(self):
        return Backup(config=self.config_name)

    def ajax_check_result(self, **kwargs):
        """AjaxTaskMixin polls request.path, which drops the ?config= the rest of the page
        carries - each poll re-enters dispatch and would rebuild the tables, and contact
        the storage, for the wrong destination."""
        path, self.request.path = self.request.path, self.request.get_full_path()
        try:
            return super().ajax_check_result(**kwargs)
        finally:
            self.request.path = path

    # noinspection PyUnresolvedReferences
    def set_cell_commands(self, table_id, row_no, html):
        self.setup_tables()
        self.add_command(overwrite_visible_cell(
            self.tables[table_id], row_no, 'Backup', f'<span class="text-success">{html}</span>')
        )

    def row_backup_schema(self, *, row_no, table_id, **_kwargs):
        self.set_cell_commands(table_id, row_no, '<div class="spinner-border spinner-border-sm"></div> Backing up')
        # the config travels as its index, the same as it does in a modal slug
        if table_id == 'schema_tables':
            if hasattr(self, 'schema'):
                # noinspection PyUnresolvedReferences
                task_kwargs = dict(config=self.config_index, schema=self.schema, table=row_no[1:])
            else:
                # noinspection PyUnresolvedReferences
                task_kwargs = dict(config=self.config_index, schema=self.kwargs['schema'], table=row_no[1:])
        else:
            # noinspection PyUnresolvedReferences
            task_kwargs = dict(config=self.config_index, schema=row_no[1:])
        return self.start_task('backup', task_kwargs=task_kwargs, result_kwargs=dict(table_id=table_id, row_no=row_no))

    def task_state_success(self, *, table_id, row_no, **_kwargs):
        self.set_cell_commands(table_id, row_no, '<i class="fas fa-check-circle"></i> Backup Complete')
        self.add_command('reload_table', table_id='files')
        return self.command_response()


class BackupContentMixin:
    """Renders the menus and datatables to a single HTML string in
    context['backup_content'] using content_template, so the backup UI can be
    embedded in a host project's own branded template. That template must still
    include the ajax_helpers/datatables/modals libs and {{ ajax_helpers_script }}."""

    content_template = 'cloud_backup/backup_content.html'

    def render_to_response(self, context, **response_kwargs):
        # with a single table (no trash on S3/Azure, one schema) DatatableView only
        # sets the singular 'datatable' key, but the content template always uses
        # 'datatables'. getattr, because the storage setup page has no tables at all
        context['datatables'] = getattr(self, 'tables', {})
        context['backup_content'] = mark_safe(
            render_to_string(self.content_template, context, request=self.request))
        return super().render_to_response(context, **response_kwargs)


class BackupBaseView(BackupContentMixin, TableBackup, PermissionRequiredMixin, MenuMixin, DatatableView):

    permission_required = 'access_admin'

    def setup_menu(self):
        self.add_config_tabs()
        self.add_menu('buttons', menu_type='buttons')
        # slugs carry the config as its index in config_names(): a modal slug is split on
        # '-', which a config name is allowed to contain (schema names are not)
        config_slug = f'config-{self.config_index}'
        if self.schema:
            self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(
                self.page_item('cloud_backup:backup_info', 'backup'),
                self.page_item('cloud_backup:schema_info', self.schema, self.schema),
            )
            self.menus['buttons'].add_items(
                (f'cloud_backup:confirm_backup,{config_slug}-schema-{self.schema}', f'BACKUP {self.schema}'),
                self.page_item('cloud_backup:schema_tables', 'View Tables', self.schema),
            )
        else:
            config = self.backup.config
            self.menus['buttons'].add_items(
                (f'cloud_backup:confirm_backup,{config_slug}', 'Backup database',
                 {'visible': config.include_db}),
                (f'cloud_backup:confirm_backup,{config_slug}-all_schemas-True', 'Backup All Schemas',
                 {'visible': config.include_db and len(self.schemas) > 1}),
                (f'cloud_backup:confirm_backup,{config_slug}-include_db-False', 'Backup Files',
                 {'visible': bool(config.dirs or config.s3_dirs)}),
                self.page_item('cloud_backup:schema_info', f'View {self.schemas[0][0]}', self.schemas[0][0],
                               visible=config.include_db and len(self.schemas) == 1),
                (f'cloud_backup:confirm_empty_trash,{config_slug}', 'Empty Trash',
                 {'css_classes': 'btn btn-warning', 'visible': self.backup.storage.supports_trash}),
                # with a single backup dir the root listing is a pointless extra
                # click, so link straight into it
                self.page_item('cloud_backup:backup_files', 'Files', 0,
                               visible=len(config.dirs) == 1),
                self.page_item('cloud_backup:backup_files_root', 'Files',
                               visible=len(config.dirs) > 1),
            )

    # noinspection PyAttributeOutsideInit
    def dispatch(self, request, *args, schema=None, **kwargs):
        self.schema = schema
        response = self.unusable_config_redirect(request, require_db=True)
        if response:
            return response
        self.schemas = get_schemas()
        return super().dispatch(request, *args, **kwargs)

    def add_tables(self):
        self.add_table('files')
        if self.backup.storage.supports_trash:
            self.add_table('deleted_files')
        if not self.schema and len(self.schemas) > 1:
            self.add_table('schemas')

    def setup_files(self, table):
        table.add_columns('.id', 'ip_address', 'table', 'name', 'size', 'encrypted',
                          # only tiered configs have anything to show here
                          DatatableColumn(column_name='tier', field='tier', enabled=self.backup.config.db_tiers),
                          DateTimeColumn(title='Backup Date', field='created'),
                          DatatableColumn(column_name='drop_restore', enabled=allowed_to_restore(),
                                          render=[row_button('drop_restore', 'Drop Restore',
                                                             button_classes='btn btn-warning btn-sm',)]))
        table.sort('-created')
        table.table_options['stateSave'] = False

    @ConfirmAjaxMethod(message='This will overwrite the current database and data could be lost')
    def row_drop_restore(self, row_data, **_kwargs):
        if not allowed_to_restore():
            return self.command_response('message', text=RESTORE_BLOCKED_MESSAGE)
        table_row = json.loads(row_data)
        return self.command_response('show_modal',
                                     modal=reverse_modal('cloud_backup:restore_db',
                                                         base64={'pk': table_row[0],
                                                                 'drop_schema': self.schema or 'public',
                                                                 'config': self.config_index}))

    @staticmethod
    def setup_deleted_files(table):
        table.add_columns('.id', 'name', 'size', DateTimeColumn(title='Backup Date', field='created'),
                          DatatableColumn(column_name='Undelete', render=[row_button(
                              'undelete', 'Undelete', button_classes='btn btn-secondary btn-sm'
                          )]))
        table.sort('-created')
        table.table_options['stateSave'] = False

    def row_undelete(self, row_no, **_kwargs):
        self.backup.storage.restore_deleted(row_no[1:])
        return self.command_response('reload')

    def setup_schemas(self, table):
        table.add_columns(
            'schema', 'size',
            ColumnLink(
                column_name='view_schema', link_ref_column='schema',
                # a url already containing the dummy id is used as it is, so the schema
                # link can keep the ?config= the rest of the page carries
                url_name=config_url('cloud_backup:schema_info', DUMMY_ID, config=self.url_config),
                link_html='<button class="btn btn-sm btn-outline-dark">VIEW</button>'
            ),
            ColumnBase(column_name='Backup',
                       render=[row_button('backup_schema', 'Backup', button_classes='btn btn-success btn-sm',)])
        )
        table.table_data = [{'schema': s[0], 'size': s[1]} for s in self.schemas]
        table.table_options['column_id'] = 0
        table.sort('schema')
        table.table_options['stateSave'] = False

    def get_context_data(self, **kwargs):
        self.add_page_command('ajax_post', data={'ajax': 'read_storage_info'})
        context = super().get_context_data(**kwargs)
        context['schema'] = self.schema
        return context

    def ajax_read_storage_info(self, **_kwargs):
        db = self.backup.get_backup_db(schema=self.schema)
        info = db.storage.storage_info(db.base_backup_dir)
        if info['web_link']:
            location = '<a target="_blank" href="{}">{}</a>'.format(info['web_link'], info['name'])
        else:
            location = info['name']
        html = f'Backup Folder {location}'
        if info['used'] is not None and info['limit'] is not None:
            gb = 1024 * 1024 * 1024
            html += ' &mdash; {:.1f} GB Used of {:.1f} GB'.format(info['used'] / gb, info['limit'] / gb)
        html = f'<div class="mb-2">{html}</div>'
        html += status_table(db.storage.protection_info(**db.tier_policy()))
        return self.command_response('html', selector='#storage_info', html=html)

    def get_table_query(self, table, **kwargs):
        files = self.backup.get_backup_db(schema=self.schema).get_db_backup_files(
            deleted=table.table_id != 'files')
        rows = [dict(**f, **f.get('metadata', {})) for f in files if not f.get('metadata', {}).get('table')]
        for row in rows:
            row['encrypted'] = 'Yes' if row.get('encrypted') else 'No'
        return rows


class BackupView(BackupBaseView):

    template_name = 'cloud_backup/backup.html'


class SchemaTableBaseView(BackupContentMixin, TableBackup, PermissionRequiredMixin, MenuMixin, DatatableView):

    permission_required = 'access_admin'

    def dispatch(self, request, *args, **kwargs):
        response = self.unusable_config_redirect(request, require_db=True)
        if response:
            return response
        return super().dispatch(request, *args, **kwargs)

    def setup_menu(self):
        self.add_config_tabs()
        self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(
            self.page_item('cloud_backup:backup_info', 'backup'),
            self.page_item('cloud_backup:schema_info', self.kwargs['schema'], self.kwargs['schema']),
            self.page_item('cloud_backup:schema_tables', 'tables', self.kwargs['schema']),
        )

    def add_tables(self):
        self.add_table('files')
        self.add_table('schema_tables')

    def setup_files(self, table):
        table.add_columns('.id', 'ip_address', 'table', 'name', 'size', 'encrypted',
                          DatatableColumn(column_name='tier', field='tier', enabled=self.backup.config.db_tiers),
                          DateTimeColumn(title='Backup Date', field='created'),
                          DatatableColumn(column_name='restore', enabled=allowed_to_restore(),
                                          render=[row_button('restore', 'Restore Table',
                                                             button_classes='btn btn-danger btn-sm')]))
        table.sort('-created')
        table.table_options['stateSave'] = False

    def row_restore(self, row_data, **_kwargs):
        """The confirm modal is opened here rather than by a ModalLink column: a datatable
        modal link builds its whole slug client-side, so the selected config has to be put
        into the base64 payload server side."""
        if not allowed_to_restore():
            return self.command_response('message', text=RESTORE_BLOCKED_MESSAGE)
        return self.command_response('show_modal',
                                     modal=reverse_modal('cloud_backup:confirm_restore_db',
                                                         base64={'pk': json.loads(row_data)[0],
                                                                 'config': self.config_index}))

    def row_download_xls(self,  **kwargs):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(get_table_column_names(self.kwargs['schema'], table_name=kwargs['row_no'][1:]))
        for r in get_table_data(self.kwargs['schema'], table_name=kwargs['row_no'][1:]):
            sheet.append([c.replace(tzinfo=None) if isinstance(c, datetime.datetime) else c for c in r])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f'{kwargs["row_no"][1:]}.xlsx'
        return self.command_response('save_file', data=base64.b64encode(output.read()).decode('ascii'),
                                     filename=filename)

    def setup_schema_tables(self, table):
        table.add_columns(
            'table', 'size', ('rows', {'title': 'No. Rows (Approx)'}),
            ColumnBase(column_name='Download',
                       render=[row_button('download_xls', '<i class="far fa-file-excel"></i>',
                                          button_classes='btn btn-outline-secondary btn-sm', )]),
            ColumnBase(column_name='Backup',
                       render=[row_button('backup_schema', 'Backup', button_classes='btn btn-success btn-sm', )])
        )
        table.table_data = [{'table': s[0], 'size': s[1], 'rows': s[2]}
                            for s in get_schema_tables(self.kwargs['schema'])]
        table.table_options['column_id'] = 0
        table.sort('table')
        table.table_options['stateSave'] = False

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['schema'] = self.kwargs['schema']
        return context

    def get_table_query(self, table, **kwargs):
        files = self.backup.get_backup_db(schema=self.kwargs.get('schema')).get_db_backup_files()
        rows = [dict(**f, **f['metadata']) for f in files if f.get('metadata', {}).get('table')]
        for row in rows:
            row['encrypted'] = 'Yes' if row.get('encrypted') else 'No'
        return rows


class SchemaTableView(SchemaTableBaseView):

    template_name = 'cloud_backup/backup.html'


class BackupFilesBaseView(BackupContentMixin, TableBackup, PermissionRequiredMixin, MenuMixin, DatatableView):
    """File browser over the BACKUP_DIRS folder backups. The root level (no backup_dir)
    lists each configured backup directory as a folder; inside one, sub-folders are
    clickable rows and files carry per-file checksum verification against the local
    source files."""

    permission_required = 'access_admin'

    # noinspection PyAttributeOutsideInit
    def dispatch(self, request, *args, backup_dir=None, sub_path='', **kwargs):
        self.backup_dir = backup_dir
        self.sub_path = sub_path.strip('/')
        response = self.unusable_config_redirect(request)
        if response:
            return response
        if backup_dir is not None:
            if not 0 <= backup_dir < len(self.backup.config.dirs):
                raise Http404('No such backup directory')
            self.source_dir, self.dest_name = self.backup.config.dirs[self.backup_dir]
        return super().dispatch(request, *args, **kwargs)

    @cached_property
    def dest_folder(self):
        # resolve the destination exactly as backup_folder does, but with get_folder
        # so browsing never creates folders
        local = BackupLocal(self.backup.storage, self.backup.config.root, self.backup.logger,
                            config=self.backup.config)
        folder = self.backup.storage.get_folder(self.dest_name, parent=local.base_backup_dir)
        for segment in self.sub_path.split('/') if self.sub_path else []:
            if folder is None:
                return None
            folder = self.backup.storage.get_folder(segment, parent=folder)
        return folder

    def setup_menu(self):
        self.add_config_tabs()
        crumbs = [self.page_item('cloud_backup:backup_info', 'backup'),
                  self.page_item('cloud_backup:backup_files_root', 'Files')]
        if self.backup_dir is not None:
            crumbs.append(self.page_item('cloud_backup:backup_files', self.dest_name, self.backup_dir))
            segments = self.sub_path.split('/') if self.sub_path else []
            for n, segment in enumerate(segments):
                crumbs.append(self.page_item('cloud_backup:backup_files_path', segment,
                                             self.backup_dir, '/'.join(segments[:n + 1])))
        self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(*crumbs)
        if self.backup_dir is not None:
            self.add_menu('buttons', menu_type='buttons').add_items(
                (f'cloud_backup:verify_files,config-{self.config_index}-backup_dir-{self.backup_dir}',
                 'Verify All Files'))

    def add_tables(self):
        self.add_table('files')

    @staticmethod
    def setup_files(table):
        table.add_columns('.id', '.path', 'name', 'size',
                          DateTimeColumn(title='Backup Date', field='created'),
                          'checksum', 'encrypted', 'verify')
        # rows arrive folders-first from get_table_query; sorting on the name column
        # would order by its HTML, so disable client-side ordering entirely
        table.table_options['ordering'] = False
        table.table_options['stateSave'] = False

    @staticmethod
    def folder_row(row_id, name, url):
        return {'id': row_id, 'path': '',
                'name': f'<a href="{url}"><i class="fas fa-folder"></i> {escape(name)}</a>',
                'size': '', 'created': None, 'checksum': '', 'encrypted': '', 'verify': ''}

    def get_table_query(self, table, **kwargs):
        if self.backup_dir is None:
            return [self.folder_row(f'dir{index}', dest_name,
                                    config_url('cloud_backup:backup_files', index, config=self.url_config))
                    for index, (_source_dir, dest_name) in enumerate(self.backup.config.dirs)]
        if self.dest_folder is None:
            return []
        rows = []
        for sub in sorted(self.backup.storage.list_folders(self.dest_folder),
                          key=lambda s: s['name'].lower()):
            sub_rel = f"{self.sub_path}/{sub['name']}" if self.sub_path else sub['name']
            rows.append(self.folder_row(
                # the id column becomes the row's DOM id, which the datatables JS looks
                # up with a jQuery selector - storage ids can be object keys whose '/'
                # and '.' break the selector, so use a digest of the path instead
                'd' + hashlib.md5(sub_rel.encode()).hexdigest(), sub['name'],
                config_url('cloud_backup:backup_files_path', self.backup_dir, sub_rel,
                           config=self.url_config)))
        verify_button = row_button('verify', 'Verify',
                                   button_classes='btn btn-outline-primary btn-sm')['html']
        # metadata is always fetched so the encrypted column is accurate and the
        # checksum column shows the plaintext md5 for files backed up while
        # encryption was on, even if it has since been turned off
        for f in sorted(self.backup.storage.list_files(self.dest_folder, include_metadata=True),
                        key=lambda s: s['name'].lower()):
            rel_path = f"{self.sub_path}/{f['name']}" if self.sub_path else f['name']
            rows.append({'id': hashlib.md5(rel_path.encode()).hexdigest(),
                         'path': rel_path,
                         'name': f'<i class="far fa-file"></i> {escape(f["name"])}',
                         'size': f['size'],
                         'created': f['created'],
                         'checksum': BackupLocal.content_hash(f) or '',
                         'encrypted': 'Yes' if (f.get('metadata') or {}).get('encrypted') else 'No',
                         'verify': verify_button})
        return rows

    def row_verify(self, row_no, row_data, **_kwargs):
        row = json.loads(row_data)
        rel_path, checksum = row[1], row[5]
        local_path = local_backup_path(self.source_dir, rel_path)
        if local_path is None or not os.path.isfile(local_path):
            badge = '<span class="badge badge-warning">Missing locally</span>'
        elif not checksum:
            badge = '<span class="badge badge-secondary">No stored checksum</span>'
        elif BackupLocal.md5sum(local_path) == checksum:
            badge = '<span class="badge badge-success"><i class="fas fa-check"></i> Match</span>'
        else:
            badge = '<span class="badge badge-danger">Changed</span>'
        self.setup_tables()
        self.add_command(overwrite_visible_cell(self.tables['files'], row_no, 'verify', badge))
        return self.command_response()

    def get_context_data(self, **kwargs):
        self.add_page_command('ajax_post', data={'ajax': 'read_storage_info'})
        return super().get_context_data(**kwargs)

    def ajax_read_storage_info(self, **_kwargs):
        if self.backup_dir is None:
            html = f'{len(self.backup.config.dirs)} backup folder(s) configured'
        elif self.dest_folder is None:
            html = f'No backups found yet for {self.source_dir}'
        else:
            info = self.backup.storage.storage_info(self.dest_folder)
            if info['web_link']:
                location = '<a target="_blank" href="{}">{}</a>'.format(info['web_link'], info['name'])
            else:
                location = info['name']
            source = self.source_dir if not self.sub_path else f'{self.source_dir}/{self.sub_path}'
            html = f'Backup Folder {location} &mdash; backed up from {source}'
        return self.command_response('html', selector='#storage_info', html=html)


class BackupFilesView(BackupFilesBaseView):

    template_name = 'cloud_backup/backup.html'


# the generated commands, kept per config so a copy button can hand back exactly what is
# on the screen. A few KB of plain text with no credentials in it
SETUP_COMMANDS_SESSION_KEY = 'cloud_backup_setup_commands'


def setup_panel(check, copy_button=None, shell_buttons=''):
    """One destination's panel: what it is, what was found, and what to run about it.
    copy_button(index=None) renders a clipboard button for one command, or for all of
    them when index is None; shell_buttons re-renders the panel for another shell."""
    # only rows that flag themselves count: most of protection_info is context, and
    # Object Lock being off is the normal state for a lifecycle-managed bucket
    actions = {row.get('action') for row in check['rows']}
    if check['error']:
        state = ('danger', 'Configuration error')
    elif check['storage_error'] or check['state'] in ('missing', 'denied') or 'fix' in actions:
        state = ('danger', 'Action required')
    elif check['state'] == 'unknown' or 'warn' in actions:
        state = ('warning', 'Worth checking')
    else:
        state = ('success', 'Ready')
    html = (f'<div class="card-body"><h5>{escape(check["name"])} '
            f'<span class="badge badge-{state[0]}">{state[1]}</span></h5>')
    if check['error']:
        return html + f'<p class="text-danger">{escape(check["error"])}</p></div>'
    html += '<table class="table table-sm w-auto"><tbody>'
    for label, value in check['facts']:
        html += f'<tr><td class="text-muted">{escape(label)}</td><td>{escape(value)}</td></tr>'
    html += '</tbody></table>'
    if check['storage_error']:
        html += (f'<p class="text-danger">The destination could not be contacted: '
                 f'{escape(check["storage_error"])}</p>')
    html += status_table(check['rows'])
    if check['guidance']:
        html += '<ul class="mt-3">' + ''.join(f'<li>{escape(line)}</li>' for line in check['guidance']) + '</ul>'
    if check['commands']:
        heading = ('Run these to set the destination up' if state[1] == 'Action required'
                   else 'Reference only - this destination already matches')
        html += f'<h6 class="mt-3">{heading} <small class="text-muted">(b2 command-line tool v4)</small></h6>'
        html += f'<div class="d-flex mb-2">{shell_buttons}<div class="ml-auto">'
        html += (copy_button() if copy_button else '') + '</div></div>'
        for index, command in enumerate(check['commands']):
            button = copy_button(index) if copy_button else ''
            html += (f'<p class="mb-1 mt-2"><small class="text-muted">{escape(command["note"])}</small></p>'
                     f'<div class="d-flex align-items-start">'
                     f'<pre class="bg-light border p-2 mb-0 flex-grow-1 text-wrap">'
                     f'{escape(command["command"])}</pre>{button}</div>')
    return html + '</div>'


class StorageSetupBaseView(BackupConfigMixin, BackupContentMixin, PermissionRequiredMixin, AjaxMenuTemplateView):
    """Read-only page describing every configured destination and the commands to set one
    up - a panel per destination, whichever one the other pages are on. Deliberately not a
    TableBackup: that resolves a config on first access, which raises when the configs are
    exactly what needs fixing. BackupConfigMixin is safe here - it only reads
    config_names()/safe_config(), never Backup()."""

    permission_required = 'access_admin'
    content_template = 'cloud_backup/storage_setup_content.html'
    setup_page = True

    def setup_menu(self):
        self.add_config_tabs()
        self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(
            (self.config_home_url(self.config_name), 'backup', {'link_type': MenuItem.HREF}),
            self.page_item('cloud_backup:storage_setup', 'storage setup'),
        )

    def get_context_data(self, **kwargs):
        names = config_names()
        for name in names:
            # one request per destination so a slow or unreachable one only holds up
            # its own panel
            self.add_page_command('ajax_post', data={'ajax': 'check_config', 'config': name})
        context = super().get_context_data(**kwargs)
        context['configs'] = [{'index': index, 'name': name} for index, name in enumerate(names)]
        return context

    def copy_button(self, config, index=None):
        """A clipboard button for one command, or for all of them when index is None.
        The copied text comes back from the server rather than being embedded in the
        link: a lifecycle rule is quoted JSON, and MenuItem's javascript hrefs swap
        double quotes for single ones, which would mangle it.

        AJAX_BUTTON cannot carry which command to copy, so the button javascript is
        built directly - the same thing AjaxButtonMenuItem does, without needing a
        version of django-tab-menus that has it."""
        kwargs = {'config': config} if index is None else {'config': config, 'index': index}
        # the button_group template supplies the btn class itself
        display = MenuItemDisplay('Copy all' if index is None else '', 'far fa-clipboard',
                                  'btn-outline-secondary btn-sm' + ('' if index is None else ' ml-2'))
        return HtmlMenu(self.request, 'button_group').add_items(
            MenuItem(button_javascript('copy_commands', **kwargs).replace('"', "'"), display,
                     link_type=MenuItem.JAVASCRIPT,
                     tooltip='Copy to the clipboard')).render()

    def shell_buttons(self, config, shell):
        """Switch the quoting the commands are written for - a JSON lifecycle rule
        survives bash, PowerShell and cmd.exe only if quoted their way."""
        return HtmlMenu(self.request, 'button_group').add_items(*[
            MenuItem(post_json_js(ajax='check_config', config=config, shell=key).replace('"', "'"),
                     MenuItemDisplay(label, css_classes='btn-outline-secondary btn-sm'
                                                        + (' active' if key == shell else '')),
                     link_type=MenuItem.JAVASCRIPT)
            for key, label in storage_setup.SHELL_LABELS.items()]).render()

    def ajax_check_config(self, config, shell=None, **_kwargs):
        names = config_names()
        if config not in names:
            return self.command_response('message', text=f'No backup configuration named {config}')
        check = storage_setup.check_config(config, shell=shell or storage_setup.shell_for_request(self.request))
        commands = [command['command'] for command in check['commands']]
        if commands:
            # cached rather than regenerated on copy: regenerating re-runs the checks and
            # could hand over text that differs from what is on the screen
            cache = self.request.session.get(SETUP_COMMANDS_SESSION_KEY, {})
            cache[config] = commands
            self.request.session[SETUP_COMMANDS_SESSION_KEY] = cache
        return self.command_response(
            'html', selector=f'#setup_config_{names.index(config)}',
            html=setup_panel(check, partial(self.copy_button, config) if commands else None,
                             shell_buttons=self.shell_buttons(config, check['shell']) if commands else ''))

    def button_copy_commands(self, config, index=None, **_kwargs):
        commands = self.request.session.get(SETUP_COMMANDS_SESSION_KEY, {}).get(config)
        if not commands:
            return self.command_response('message', text='Reload the page and try again')
        if index is None:
            return self.command_response('clipboard', text='\n'.join(commands))
        try:
            command = commands[int(index)]
        except (TypeError, ValueError, IndexError):
            return self.command_response('message', text='Reload the page and try again')
        return self.command_response('clipboard', text=command)


class StorageSetupView(StorageSetupBaseView):

    template_name = 'cloud_backup/backup.html'
