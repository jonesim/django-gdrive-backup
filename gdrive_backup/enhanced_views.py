import base64
import datetime
import json
from functools import cached_property
from io import BytesIO

from ajax_helpers.mixins import AjaxHelpers, AjaxTaskMixin
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.template.loader import render_to_string
from django.utils.safestring import mark_safe
from django_datatables.columns import DateTimeColumn, DatatableColumn, ColumnLink, ColumnBase
from django_datatables.datatables import DatatableView
from django_datatables.helpers import row_button, overwrite_cell
from django_menus.menu import MenuMixin
from django_modals.datatables import ModalLink
from django_modals.decorators import ConfirmAjaxMethod
from django_modals.helper import reverse_modal
from openpyxl import Workbook

from gdrive_backup.backup import Backup
from .sql_functions import get_schemas, get_schema_tables, get_table_column_names, get_table_data
from .tasks import ajax_backup
from .utils import allowed_to_restore, RESTORE_BLOCKED_MESSAGE


def restore_table_button(text):
    return ModalLink(row=True, base64=True, modal_name='gdrive_backup:confirm_restore_db',
                     css_class='btn btn-danger btn-sm', title='Restore', button_text=text,
                     enabled=allowed_to_restore())


class TableBackup(AjaxTaskMixin, AjaxHelpers):

    tasks = {'backup': ajax_backup}

    @cached_property
    def backup(self):
        return Backup()

    # noinspection PyUnresolvedReferences
    def set_cell_commands(self, table_id, row_no, html):
        self.setup_tables()
        self.add_command(overwrite_cell(
            self.tables[table_id], row_no, 'Backup', f'<span class="text-success">{html}</span>')
        )

    def row_backup_schema(self, *, row_no, table_id, **_kwargs):
        self.set_cell_commands(table_id, row_no, '<div class="spinner-border spinner-border-sm"></div> Backing up')
        if table_id == 'schema_tables':
            if hasattr(self, 'schema'):
                # noinspection PyUnresolvedReferences
                task_kwargs = dict(schema=self.schema, table=row_no[1:])
            else:
                # noinspection PyUnresolvedReferences
                task_kwargs = dict(schema=self.kwargs['schema'], table=row_no[1:])
        else:
            # noinspection PyUnresolvedReferences
            task_kwargs = dict(schema=row_no[1:])
        return self.start_task('backup', task_kwargs=task_kwargs, result_kwargs=dict(table_id=table_id, row_no=row_no))

    def task_state_success(self, *, table_id, row_no, **_kwargs):
        self.set_cell_commands(table_id, row_no, '<i class="fas fa-check-circle"></i> Backup Complete')
        self.add_command('reload_table', table_id='files')
        return self.command_response()


class BackupContentMixin:
    """Renders the menus and datatables to a single HTML string in
    context['backup_content'] using content_template, so the backup UI can be
    embedded in a host project's own branded template. That template must still
    include the ajax_helpers/datatables/modals libs and {{ ajax_helpers_script }}."""

    content_template = 'gdrive_backup/backup_content.html'

    def render_to_response(self, context, **response_kwargs):
        # with a single table (no trash on S3/Azure, one schema) DatatableView only
        # sets the singular 'datatable' key, but the content template always uses
        # 'datatables'
        context['datatables'] = self.tables
        context['backup_content'] = mark_safe(
            render_to_string(self.content_template, context, request=self.request))
        return super().render_to_response(context, **response_kwargs)


class BackupBaseView(BackupContentMixin, TableBackup, PermissionRequiredMixin, MenuMixin, DatatableView):

    permission_required = 'access_admin'

    def setup_menu(self):
        self.add_menu('buttons', menu_type='buttons')
        if self.schema:
            self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(
                ('gdrive_backup:backup_info', 'backup'),
                ('gdrive_backup:schema_info', self.schema, {'url_args': [self.schema]}),
            )
            self.menus['buttons'].add_items(
                (f'gdrive_backup:confirm_backup,schema-{self.schema}', f'BACKUP {self.schema}'),
                ('gdrive_backup:schema_tables', 'View Tables', {'url_args': [self.schema]}),
            )
        else:
            self.menus['buttons'].add_items(
                ('gdrive_backup:confirm_backup,-', 'Backup database'),
                ('gdrive_backup:confirm_backup,all_schemas-True', 'Backup All Schemas',
                 {'visible': len(self.schemas) > 1}),
                (f'gdrive_backup:schema_info,{self.schemas[0][0]}', f'View {self.schemas[0][0]}',
                 {'visible': len(self.schemas) == 1}),
                ('gdrive_backup:confirm_empty_trash', 'Empty Trash',
                 {'css_classes': 'btn btn-warning', 'visible': self.backup.storage.supports_trash}),
                ('gdrive_backup:confirm_drop_schema,-', 'Drop Public Schema',
                 {'css_classes': 'btn btn-danger', 'visible': allowed_to_restore()}),
            )

    # noinspection PyAttributeOutsideInit
    def dispatch(self, request, *args, schema=None, **kwargs):
        self.schema = schema
        self.schemas = get_schemas()
        return super().dispatch(request, *args, **kwargs)

    def add_tables(self):
        self.add_table('files')
        if self.backup.storage.supports_trash:
            self.add_table('deleted_files')
        if not self.schema and len(self.schemas) > 1:
            self.add_table('schemas')

    @staticmethod
    def setup_files(table):
        table.add_columns('.id', 'ip_address', 'table', 'name', 'size',
                          DateTimeColumn(title='Backup Date', field='created'),
                          DatatableColumn(column_name='drop_restore', enabled=allowed_to_restore(),
                                          render=[row_button('drop_restore', 'Drop Restore',
                                                             button_classes='btn btn-warning btn-sm',)]),
                          restore_table_button('Restore DB'))
        table.sort('-created')
        table.table_options['stateSave'] = False

    @ConfirmAjaxMethod(message='This will overwrite the current database and data could be lost')
    def row_drop_restore(self, row_data, **_kwargs):
        if not allowed_to_restore():
            return self.command_response('message', text=RESTORE_BLOCKED_MESSAGE)
        table_row = json.loads(row_data)
        return self.command_response('show_modal',
                                     modal=reverse_modal('gdrive_backup:restore_db'
                                                         ,base64={'pk': table_row[0], 'drop_schema': 'public'}))

    @staticmethod
    def setup_deleted_files(table):
        table.add_columns('.id', 'name', 'size', DateTimeColumn(title='Backup Date', field='created'),
                          DatatableColumn(column_name='Undelete', render=[row_button(
                              'undelete', 'Undelete', button_classes='btn btn-secondary btn-sm'
                          )]))
        table.sort('-created')
        table.table_options['stateSave'] = False

    def row_undelete(self, row_no, **_kwargs):
        self.backup.storage.restore_deleted(row_no[1:])
        return self.command_response('reload')

    def setup_schemas(self, table):
        table.add_columns(
            'schema', 'size',
            ColumnLink(
                column_name='view_schema', link_ref_column='schema', url_name='gdrive_backup:schema_info',
                link_html='<button class="btn btn-sm btn-outline-dark">VIEW</button>'
            ),
            ColumnBase(column_name='Backup',
                       render=[row_button('backup_schema', 'Backup', button_classes='btn btn-success btn-sm',)])
        )
        table.table_data = [{'schema': s[0], 'size': s[1]} for s in self.schemas]
        table.table_options['column_id'] = 0
        table.sort('schema')
        table.table_options['stateSave'] = False

    def get_context_data(self, **kwargs):
        self.add_page_command('ajax_post', data={'ajax': 'read_storage_info'})
        context = super().get_context_data(**kwargs)
        context['schema'] = self.schema
        return context

    def ajax_read_storage_info(self, **_kwargs):
        db = self.backup.get_backup_db(schema=self.schema)
        info = db.storage.storage_info(db.base_backup_dir)
        if info['web_link']:
            location = '<a target="_blank" href="{}">{}</a>'.format(info['web_link'], info['name'])
        else:
            location = info['name']
        html = f'Backup Folder {location}'
        if info['used'] is not None and info['limit'] is not None:
            gb = 1024 * 1024 * 1024
            html += '<br>{:.1f} GB Used of {:.1f} GB'.format(info['used'] / gb, info['limit'] / gb)
        badge_colours = {'Enabled': 'success', 'Disabled': 'danger', 'Suspended': 'warning'}
        for p in db.storage.protection_info():
            colour = badge_colours.get(p['status'], 'secondary')
            html += f'<br>{p["label"]} <span class="badge badge-{colour}">{p["status"]}</span>'
            if p.get('detail'):
                html += f' <small class="text-muted">{p["detail"]}</small>'
        return self.command_response('html', selector='#storage_info', html=html)

    def get_table_query(self, table, **kwargs):
        files = self.backup.get_backup_db(schema=self.schema).get_db_backup_files(
            deleted=table.table_id != 'files')
        return [dict(**f, **f.get('metadata', {})) for f in files if not f.get('metadata', {}).get('table')]


class BackupView(BackupBaseView):

    template_name = 'gdrive_backup/backup.html'


class SchemaTableBaseView(BackupContentMixin, TableBackup, PermissionRequiredMixin, MenuMixin, DatatableView):

    permission_required = 'access_admin'

    def setup_menu(self):
        self.add_menu('breadcrumbs', menu_type='breadcrumb').add_items(
            ('gdrive_backup:backup_info', 'backup'),
            ('gdrive_backup:schema_info', self.kwargs['schema'], {'url_args': [self.kwargs['schema']]}),
            ('gdrive_backup:schema_tables', 'tables', {'url_args': [self.kwargs['schema']]}),
        )

    def add_tables(self):
        self.add_table('files')
        self.add_table('schema_tables')

    @staticmethod
    def setup_files(table):
        table.add_columns('.id', 'ip_address', 'table', 'name', 'size',
                          DateTimeColumn(title='Backup Date', field='created'),
                          restore_table_button('Restore Table'))
        table.sort('-created')
        table.table_options['stateSave'] = False

    def row_download_xls(self,  **kwargs):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(get_table_column_names(self.kwargs['schema'], table_name=kwargs['row_no'][1:]))
        for r in get_table_data(self.kwargs['schema'], table_name=kwargs['row_no'][1:]):
            sheet.append([c.replace(tzinfo=None) if isinstance(c, datetime.datetime) else c for c in r])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f'{kwargs["row_no"][1:]}.xlsx'
        return self.command_response('save_file', data=base64.b64encode(output.read()).decode('ascii'),
                                     filename=filename)

    def setup_schema_tables(self, table):
        table.add_columns(
            'table', 'size', ('rows', {'title': 'No. Rows (Approx)'}),
            ColumnBase(column_name='Download',
                       render=[row_button('download_xls', '<i class="far fa-file-excel"></i>',
                                          button_classes='btn btn-outline-secondary btn-sm', )]),
            ColumnBase(column_name='Backup',
                       render=[row_button('backup_schema', 'Backup', button_classes='btn btn-success btn-sm', )])
        )
        table.table_data = [{'table': s[0], 'size': s[1], 'rows': s[2]}
                            for s in get_schema_tables(self.kwargs['schema'])]
        table.table_options['column_id'] = 0
        table.sort('table')
        table.table_options['stateSave'] = False

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['schema'] = self.kwargs['schema']
        return context

    def get_table_query(self, table, **kwargs):
        files = self.backup.get_backup_db(schema=self.kwargs.get('schema')).get_db_backup_files()
        return [dict(**f, **f['metadata']) for f in files if f.get('metadata', {}).get('table')]


class SchemaTableView(SchemaTableBaseView):

    template_name = 'gdrive_backup/backup.html'
